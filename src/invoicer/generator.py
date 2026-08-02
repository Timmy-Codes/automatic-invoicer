from docx import Document


class InvoiceGenerator:
    """
    Generate invoices from a Word template.

    Args:
        template_path: Path to the DOCX invoice template.
    """
    def __init__(self, template_path: str):
        self.template_path = template_path
        self.document = None

    def load_template(self) -> None:
        """Load the configured DOCX Word template into memory."""
        self.document = Document(self.template_path)

    def save(self, output_path: str) -> None:
        """
        Saves the generated DOCX invoice to a specified location.

        Args:
           output_path: Path to save the DOCX generated invoice.
        Raises:
            RuntimeError: if no template is loaded.
        """
        if self.document is None:
            raise RuntimeError("No document loaded. Call load_template() before saving.")

        self.document.save(output_path)

    def replace_placeholders(self, invoice_data: dict[str,str]) -> None:
        """
        Replace placeholders in the loaded template with actual invoice data.

        Args:
            invoice_data: A dictionary containing placeholder keys and their corresponding values.
        Raises:
            RuntimeError: if no template is loaded.
        """
        if self.document is None:
            raise RuntimeError("No document loaded. Call load_template() before replacing placeholders.")

        for paragraph in self.document.paragraphs:
            for key, value in invoice_data.items():
                placeholder = "{{" + key + "}}"
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, value)